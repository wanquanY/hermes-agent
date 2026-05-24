#!/usr/bin/env python3
"""Doxie-owned document parsing tool for Hermes local attachments."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import mimetypes
import os
import re
import shutil
import subprocess
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from agent.file_safety import get_read_block_error
from tools.registry import registry, tool_error

logger = logging.getLogger(__name__)

try:
    from tools.file_tools import _resolve_path_for_task
except Exception:  # pragma: no cover
    def _resolve_path_for_task(filepath: str, task_id: str = "default") -> Path:
        p = Path(filepath).expanduser()
        if not p.is_absolute():
            p = Path(os.getcwd()) / p
        return p.resolve()


DEFAULT_MAX_CHARS = 60_000
MAX_ALLOWED_CHARS = 120_000
MINERU_SUPPORTED_EXTENSIONS = {"pdf", "ppt", "pptx"}
MINERU_DEFAULT_MODEL_VERSION = "vlm"
DOXIE_MINERU_PROXY_TIMEOUT_SECONDS = 600
SUPPORTED_EXTENSIONS = {
    "pdf", "docx", "doc", "pptx", "ppt", "xlsx", "xls", "csv", "tsv",
    "txt", "md", "markdown", "json", "xml", "yaml", "yml", "toml", "ini",
    "log", "rtf",
}
TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "json", "xml", "yaml", "yml", "toml", "ini",
    "log", "csv", "tsv", "rtf",
    "py", "js", "jsx", "ts", "tsx", "vue", "html", "css", "scss", "sass",
    "less", "java", "kt", "swift", "go", "rs", "c", "cc", "cpp", "h", "hpp",
    "cs", "php", "rb", "sh", "bash", "zsh", "fish", "sql", "graphql",
}
TEXT_MIME_TYPES = {
    "application/json",
    "application/xml",
    "application/x-yaml",
    "application/yaml",
    "application/toml",
    "application/rtf",
}
ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1")


class MinerUError(RuntimeError):
    """Raised when the Doxie MinerU proxy is unavailable or returns an error."""


def _safe_int(value: Any, default: int, *, minimum: int = 0, maximum: int | None = None) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        parsed = default
    parsed = max(minimum, parsed)
    if maximum is not None:
        parsed = min(maximum, parsed)
    return parsed


def _file_type(path: Path, explicit_type: str = "") -> str:
    value = str(explicit_type or "").strip().lower().lstrip(".")
    if value:
        return value.split(";")[0].split("/")[-1]
    return path.suffix.lower().lstrip(".")


def _mime_type(path: Path, explicit_mime: str = "") -> str:
    return explicit_mime or mimetypes.guess_type(str(path))[0] or ""


def _is_text_document(ext: str, mime_type: str) -> bool:
    normalized_mime = mime_type.lower()
    return ext in TEXT_EXTENSIONS or normalized_mime.startswith("text/") or normalized_mime in TEXT_MIME_TYPES


def _env_int(name: str, default: int) -> int:
    try:
        value = int(str(os.getenv(name, "")).strip())
    except (TypeError, ValueError):
        value = default
    return value if value > 0 else default


def _multipart_body(fields: dict[str, str], files: dict[str, tuple[str, bytes, str]]) -> tuple[bytes, str]:
    boundary = f"doxie-hermes-{hashlib.sha256(os.urandom(16)).hexdigest()}"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.append(f"--{boundary}\r\n".encode())
        chunks.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        chunks.append(str(value).encode("utf-8"))
        chunks.append(b"\r\n")
    for name, (filename, data, content_type) in files.items():
        safe_filename = str(filename or "document").replace('"', '\\"')
        chunks.append(f"--{boundary}\r\n".encode())
        chunks.append(
            f'Content-Disposition: form-data; name="{name}"; filename="{safe_filename}"\r\n'.encode()
        )
        chunks.append(f"Content-Type: {content_type or 'application/octet-stream'}\r\n\r\n".encode())
        chunks.append(data)
        chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def _http_post_multipart(
    url: str,
    *,
    token: str,
    fields: dict[str, str],
    files: dict[str, tuple[str, bytes, str]],
    timeout: int,
) -> dict[str, Any]:
    body, content_type = _multipart_body(fields, files)
    request = urllib.request.Request(
        url,
        data=body,
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "Content-Type": content_type,
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read()
            status = getattr(response, "status", 200)
    except urllib.error.HTTPError as exc:
        raw = exc.read()
        status = exc.code
    except urllib.error.URLError as exc:
        raise MinerUError(f"Doxie MinerU proxy request failed: {exc}") from exc

    try:
        result = json.loads(raw.decode("utf-8"))
    except Exception as exc:
        raise MinerUError(f"Doxie MinerU proxy returned non-JSON response: HTTP {status}") from exc
    if status < 200 or status >= 300 or result.get("success") is False:
        detail = result.get("detail") if isinstance(result, dict) else None
        message = detail.get("message") if isinstance(detail, dict) else detail
        raise MinerUError(str(message or result.get("message") or f"HTTP {status}"))
    return result


def _doxie_mineru_proxy_url() -> str:
    return str(os.getenv("DOXIE_MINERU_PROXY_URL") or "").strip()


def _doxie_runtime_token() -> str:
    return str(os.getenv("DOXIE_LLM_RUNTIME_TOKEN") or "").strip()


def _parse_with_doxie_mineru_proxy(
    *,
    path: Path | None,
    file_url: str,
    file_type: str,
    mime_type: str,
) -> dict[str, Any]:
    proxy_url = _doxie_mineru_proxy_url()
    token = _doxie_runtime_token()
    if not proxy_url:
        raise MinerUError("Doxie MinerU proxy is not configured")
    if not token:
        raise MinerUError("Doxie runtime token is not configured")

    fields = {
        "file_type": file_type,
        "model_version": str(os.getenv("MINERU_MODEL_VERSION") or MINERU_DEFAULT_MODEL_VERSION).strip(),
    }
    files: dict[str, tuple[str, bytes, str]] = {}
    normalized_url = str(file_url or "").strip()
    if normalized_url:
        fields["file_url"] = normalized_url
        fields["file_name"] = Path(urllib.parse.urlparse(normalized_url).path).name or "document"
    else:
        if path is None:
            raise MinerUError("Doxie MinerU proxy local parsing requires 'path'")
        fields["file_name"] = path.name
        files["file"] = (path.name, path.read_bytes(), mime_type or "application/octet-stream")

    result = _http_post_multipart(
        proxy_url,
        token=token,
        fields={key: value for key, value in fields.items() if value},
        files=files,
        timeout=_env_int("DOXIE_MINERU_PROXY_TIMEOUT", DOXIE_MINERU_PROXY_TIMEOUT_SECONDS),
    )
    return {
        "content": str(result.get("content") or ""),
        "parser": str(result.get("parser") or "mineru"),
        "metadata": {
            **(result.get("metadata") or {}),
            "source": "doxie_mineru_proxy",
        },
    }


def _decode_bytes(data: bytes) -> tuple[str, str]:
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def _read_text(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    content, encoding = _decode_bytes(raw)
    return {
        "content": content,
        "parser": "direct_text",
        "metadata": {
            "encoding": encoding,
            "lines": len(content.splitlines()),
            "characters": len(content),
            "bytes": len(raw),
        },
    }


def _format_csv(path: Path, delimiter: str) -> dict[str, Any]:
    text_result = _read_text(path)
    rows: list[str] = []
    reader = csv.reader(text_result["content"].splitlines(), delimiter=delimiter)
    row_count = 0
    column_count = 0
    for row in reader:
        row_count += 1
        column_count = max(column_count, len(row))
        rows.append("\t".join(row))
    return {
        "content": "\n".join(rows),
        "parser": "csv",
        "metadata": {
            **text_result["metadata"],
            "rows": row_count,
            "columns": column_count,
            "delimiter": delimiter,
        },
    }


def _xml_texts(element: ET.Element, tag_suffix: str) -> list[str]:
    suffix = "}" + tag_suffix
    return [
        child.text or ""
        for child in element.iter()
        if child.tag == tag_suffix or child.tag.endswith(suffix)
        if child.text
    ]


def _parse_docx_xml(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path, "r") as zf:
        xml_bytes = zf.read("word/document.xml")

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.fromstring(xml_bytes)
    body = root.find(".//w:body", ns)
    if body is None:
        return {"content": "", "parser": "docx_xml", "metadata": {"paragraphs": 0, "tables": 0}}

    parts: list[str] = []
    paragraph_count = 0
    table_count = 0
    for paragraph in body.findall("w:p", ns):
        text = "".join(_xml_texts(paragraph, "t")).strip()
        if text:
            parts.append(text)
            paragraph_count += 1
    for table in body.findall("w:tbl", ns):
        table_count += 1
        lines: list[str] = []
        for row in table.findall(".//w:tr", ns):
            cells = [
                " ".join(t.strip() for t in _xml_texts(cell, "t") if t.strip())
                for cell in row.findall(".//w:tc", ns)
            ]
            if any(cells):
                lines.append("\t".join(cells))
        if lines:
            parts.append("[Table]\n" + "\n".join(lines))
    return {
        "content": "\n\n".join(parts).strip(),
        "parser": "docx_xml",
        "metadata": {"paragraphs": paragraph_count, "tables": table_count},
    }


def _parse_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document  # type: ignore

        document = Document(str(path))
        parts: list[str] = []
        paragraphs = 0
        tables = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
                paragraphs += 1
        for table in document.tables:
            tables += 1
            rows = [
                "\t".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            rows = [row for row in rows if row.strip()]
            if rows:
                parts.append("[Table]\n" + "\n".join(rows))
        return {
            "content": "\n\n".join(parts).strip(),
            "parser": "python_docx",
            "metadata": {"paragraphs": paragraphs, "tables": tables},
        }
    except Exception:
        return _parse_docx_xml(path)


def _parse_pptx(path: Path) -> dict[str, Any]:
    slide_pattern = re.compile(r"ppt/slides/slide(\d+)\.xml$")
    with zipfile.ZipFile(path, "r") as zf:
        slide_names = sorted(
            (name for name in zf.namelist() if slide_pattern.match(name)),
            key=lambda name: int(slide_pattern.match(name).group(1)),  # type: ignore[union-attr]
        )
        parts: list[str] = []
        for index, name in enumerate(slide_names, 1):
            root = ET.fromstring(zf.read(name))
            texts = [text.strip() for text in _xml_texts(root, "t") if text.strip()]
            if texts:
                parts.append(f"[Slide {index}]\n" + "\n".join(texts))
    return {
        "content": "\n\n".join(parts).strip(),
        "parser": "pptx_xml",
        "metadata": {"slides": len(slide_names)},
    }


def _parse_pdf_with_pdfplumber(path: Path) -> dict[str, Any]:
    import pdfplumber  # type: ignore

    content_parts: list[str] = []
    table_count = 0
    with pdfplumber.open(str(path)) as pdf:
        for index, page in enumerate(pdf.pages, 1):
            page_parts: list[str] = []
            text = page.extract_text() or ""
            if text.strip():
                page_parts.append(text.strip())
            for table in page.extract_tables() or []:
                table_count += 1
                rows = [
                    "\t".join(cell or "" for cell in row)
                    for row in table
                    if any(cell for cell in row)
                ]
                if rows:
                    page_parts.append("[Table]\n" + "\n".join(rows))
            if page_parts:
                content_parts.append(f"[Page {index}]\n" + "\n\n".join(page_parts))
    return {
        "content": "\n\n".join(content_parts).strip(),
        "parser": "pdfplumber",
        "metadata": {"pages": len(pdf.pages), "tables": table_count},
    }


def _parse_pdf_with_pypdf(path: Path) -> dict[str, Any]:
    from PyPDF2 import PdfReader  # type: ignore

    reader = PdfReader(str(path))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[Page {index}]\n{text.strip()}")
    return {
        "content": "\n\n".join(parts).strip(),
        "parser": "pypdf2",
        "metadata": {"pages": len(reader.pages), "tables": 0},
    }


def _parse_pdf_with_pdftotext(path: Path) -> dict[str, Any]:
    binary = shutil.which("pdftotext")
    if not binary:
        raise RuntimeError("No PDF parser is available. Install pdfplumber/PyPDF2 or pdftotext.")
    completed = subprocess.run(
        [binary, "-layout", str(path), "-"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=60,
    )
    return {
        "content": completed.stdout.strip(),
        "parser": "pdftotext",
        "metadata": {"pages": None, "tables": None},
    }


def _parse_pdf(path: Path) -> dict[str, Any]:
    errors: list[str] = []
    for parser in (_parse_pdf_with_pdfplumber, _parse_pdf_with_pypdf, _parse_pdf_with_pdftotext):
        try:
            return parser(path)
        except Exception as exc:
            errors.append(f"{parser.__name__}: {exc}")
    raise RuntimeError("; ".join(errors))


def _parse_xlsx(path: Path) -> dict[str, Any]:
    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        raise RuntimeError("XLSX parsing requires openpyxl in the Hermes runtime.") from exc

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    total_rows = 0
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            if not any(cell is not None for cell in row):
                continue
            rows.append("\t".join("" if cell is None else str(cell) for cell in row))
            total_rows += 1
        if rows:
            parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    return {
        "content": "\n\n".join(parts).strip(),
        "parser": "openpyxl",
        "metadata": {"sheets": len(workbook.sheetnames), "total_rows": total_rows},
    }


def _parse_doc(path: Path) -> dict[str, Any]:
    try:
        import docx2txt  # type: ignore
    except Exception as exc:
        raise RuntimeError("Legacy .doc parsing requires docx2txt in the Hermes runtime.") from exc
    return {
        "content": str(docx2txt.process(str(path)) or "").strip(),
        "parser": "docx2txt",
        "metadata": {"format": "doc"},
    }


def _parse_xls(path: Path) -> dict[str, Any]:
    try:
        import pandas as pd  # type: ignore
    except Exception as exc:
        raise RuntimeError("Legacy .xls parsing requires pandas/xlrd in the Hermes runtime.") from exc

    excel = pd.ExcelFile(str(path))
    parts: list[str] = []
    total_rows = 0
    for sheet_name in excel.sheet_names:
        frame = pd.read_excel(str(path), sheet_name=sheet_name)
        total_rows += len(frame)
        parts.append(f"[Sheet: {sheet_name}]\n{frame.to_string(index=False)}")
    return {
        "content": "\n\n".join(parts).strip(),
        "parser": "pandas",
        "metadata": {"sheets": len(excel.sheet_names), "total_rows": total_rows},
    }


def _parse_by_type(path: Path, ext: str, mime_type: str) -> dict[str, Any]:
    if _is_text_document(ext, mime_type):
        if ext == "csv":
            return _format_csv(path, ",")
        if ext == "tsv":
            return _format_csv(path, "\t")
        return _read_text(path)
    if ext == "pdf":
        return _parse_pdf(path)
    if ext == "docx":
        return _parse_docx(path)
    if ext == "doc":
        return _parse_doc(path)
    if ext == "pptx":
        return _parse_pptx(path)
    if ext == "ppt":
        raise RuntimeError("Legacy .ppt parsing is not supported locally. Convert it to .pptx or PDF first.")
    if ext == "xlsx":
        return _parse_xlsx(path)
    if ext == "xls":
        return _parse_xls(path)
    raise RuntimeError(f"Unsupported document type: {ext or mime_type or 'unknown'}")


def _slice_content(content: str, offset: int, max_chars: int) -> dict[str, Any]:
    start = min(offset, len(content))
    end = min(start + max_chars, len(content))
    return {
        "content": content[start:end],
        "offset": start,
        "max_chars": max_chars,
        "next_offset": end if end < len(content) else None,
        "truncated": end < len(content),
        "total_characters": len(content),
    }


def parse_document_tool(
    path: str = "",
    file_url: str = "",
    file_type: str = "",
    mime_type: str = "",
    offset: int = 0,
    max_chars: int = DEFAULT_MAX_CHARS,
    allow_local_fallback: bool = True,
    task_id: str = "default",
) -> str:
    normalized_path = str(path or "").strip()
    normalized_file_url = str(file_url or "").strip()
    if not normalized_path and not normalized_file_url:
        return tool_error("parse_document: missing required field 'path' or 'file_url'.")

    resolved: Path | None = None
    if normalized_path:
        try:
            resolved = _resolve_path_for_task(normalized_path, task_id)
        except Exception as exc:
            return tool_error(f"parse_document: invalid path: {exc}")

        literal_block = get_read_block_error(normalized_path)
        resolved_block = get_read_block_error(str(resolved))
        if literal_block or resolved_block:
            return tool_error(literal_block or resolved_block or "parse_document: blocked path.")
        if not resolved.exists():
            return tool_error(f"parse_document: file not found: {resolved}")
        if not resolved.is_file():
            return tool_error(f"parse_document: path is not a file: {resolved}")

    url_path = Path(urllib.parse.urlparse(normalized_file_url).path) if normalized_file_url else Path("")
    ext = _file_type(resolved or url_path, file_type)
    guessed_mime = _mime_type(resolved or url_path, mime_type)
    if ext not in SUPPORTED_EXTENSIONS and not _is_text_document(ext, guessed_mime):
        return tool_error(
            "parse_document: unsupported file type "
            f"'{ext or guessed_mime or 'unknown'}'. Supported types: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    safe_offset = _safe_int(offset, 0, minimum=0)
    safe_max_chars = _safe_int(max_chars, DEFAULT_MAX_CHARS, minimum=1, maximum=MAX_ALLOWED_CHARS)

    mineru_error = ""
    if ext in MINERU_SUPPORTED_EXTENSIONS:
        try:
            parsed = _parse_with_doxie_mineru_proxy(
                path=resolved,
                file_url=normalized_file_url,
                file_type=ext,
                mime_type=guessed_mime,
            )
        except Exception as exc:
            mineru_error = str(exc)
            logger.warning(
                "parse_document: Doxie MinerU proxy failed file_type=%s file_name=%s fallback=%s error=%s",
                ext,
                (resolved.name if resolved else url_path.name),
                allow_local_fallback,
                mineru_error,
            )
            if not allow_local_fallback or resolved is None:
                return tool_error(f"parse_document: MinerU failed for {ext}: {mineru_error}")
        else:
            content = str(parsed.get("content") or "")
            sliced = _slice_content(content, safe_offset, safe_max_chars)
            result = {
                "success": True,
                "path": str(resolved) if resolved else None,
                "file_url": normalized_file_url,
                "file_name": (resolved.name if resolved else url_path.name),
                "file_type": ext,
                "mime_type": guessed_mime or None,
                "parser": parsed.get("parser") or "mineru",
                "content": sliced["content"],
                "metadata": {
                    **(parsed.get("metadata") or {}),
                    "source": "mineru",
                    "file_url": normalized_file_url,
                    "local_upload": bool(resolved and not normalized_file_url),
                },
                "pagination": {
                    "offset": sliced["offset"],
                    "max_chars": sliced["max_chars"],
                    "next_offset": sliced["next_offset"],
                    "truncated": sliced["truncated"],
                    "total_characters": sliced["total_characters"],
                },
            }
            if sliced["truncated"]:
                result["hint"] = "Call parse_document again with pagination.next_offset to continue reading."
            return json.dumps(result, ensure_ascii=False)

    if resolved is None:
        return tool_error(
            "parse_document: local fallback requires 'path'. "
            f"MinerU was not used for file type '{ext}'"
            + (f": {mineru_error}" if mineru_error else ".")
        )

    try:
        parsed = _parse_by_type(resolved, ext, guessed_mime)
    except Exception as exc:
        return tool_error(f"parse_document: failed to parse {resolved.name}: {exc}")

    content = str(parsed.get("content") or "")
    sliced = _slice_content(content, safe_offset, safe_max_chars)
    result = {
        "success": True,
        "path": str(resolved),
        "file_url": normalized_file_url or None,
        "file_name": resolved.name,
        "file_type": ext,
        "mime_type": guessed_mime or None,
        "parser": parsed.get("parser") or "unknown",
        "content": sliced["content"],
        "metadata": {
            **(parsed.get("metadata") or {}),
            **({"mineru_fallback_reason": mineru_error} if mineru_error else {}),
        },
        "pagination": {
            "offset": sliced["offset"],
            "max_chars": sliced["max_chars"],
            "next_offset": sliced["next_offset"],
            "truncated": sliced["truncated"],
            "total_characters": sliced["total_characters"],
        },
    }
    if sliced["truncated"]:
        result["hint"] = "Call parse_document again with pagination.next_offset to continue reading."
    return json.dumps(result, ensure_ascii=False)


DOCUMENT_PARSE_SCHEMA = {
    "name": "parse_document",
    "description": (
        "Parse an uploaded or local document into text. Use this for PDFs, "
        "PowerPoint, Word, Excel, CSV, JSON, XML, Markdown, and TXT attachments "
        "before answering questions about their contents. PDF/PPT/PPTX documents "
        "use the Doxie-managed MinerU proxy when available, with safe local "
        "fallback for supported formats."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Absolute, relative, or ~/ path to the local document attachment.",
            },
            "file_url": {
                "type": "string",
                "description": "Optional cloud/public HTTP(S) URL for MinerU parsing.",
            },
            "file_type": {
                "type": "string",
                "description": "Optional file extension override, e.g. pdf, docx, pptx, xlsx, txt.",
            },
            "mime_type": {
                "type": "string",
                "description": "Optional MIME type override from attachment metadata.",
            },
            "offset": {
                "type": "integer",
                "description": "Character offset for paginating extracted text.",
                "default": 0,
                "minimum": 0,
            },
            "max_chars": {
                "type": "integer",
                "description": f"Maximum extracted characters to return. Default: {DEFAULT_MAX_CHARS}, max: {MAX_ALLOWED_CHARS}.",
                "default": DEFAULT_MAX_CHARS,
                "minimum": 1,
                "maximum": MAX_ALLOWED_CHARS,
            },
            "allow_local_fallback": {
                "type": "boolean",
                "description": "Allow local parser fallback if MinerU fails or is unavailable.",
                "default": True,
            },
        },
        "required": [],
    },
}


def _handle_parse_document(args: dict[str, Any], **kw: Any) -> str:
    return parse_document_tool(
        path=args.get("path", ""),
        file_url=args.get("file_url", "") or args.get("fileUrl", ""),
        file_type=args.get("file_type", ""),
        mime_type=args.get("mime_type", ""),
        offset=args.get("offset", 0),
        max_chars=args.get("max_chars", DEFAULT_MAX_CHARS),
        allow_local_fallback=bool(args.get("allow_local_fallback", True)),
        task_id=kw.get("task_id") or "default",
    )


def register_document_parse_tool() -> None:
    registry.register(
        name="parse_document",
        toolset="file",
        schema=DOCUMENT_PARSE_SCHEMA,
        handler=_handle_parse_document,
        check_fn=lambda: True,
        emoji="📄",
        max_result_size_chars=MAX_ALLOWED_CHARS,
        override=True,
    )


register_document_parse_tool()
